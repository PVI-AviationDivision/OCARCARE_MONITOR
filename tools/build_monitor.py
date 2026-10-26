# -*- coding: utf-8 -*-
"""
OCARCARE MONITOR — báo cáo hằng ngày OCARCARE (VPBank × OPES × BH PVI), luỹ kế từ 21/09/2026.

Đọc dữ liệu trong thư mục DuLieu (cùng cấp với thư mục Monitor):
    DuLieu/input/          báo cáo BH PVI hằng ngày (lấy file MỚI NHẤT — số liệu luỹ kế)
    DuLieu/mapping/        file mapping điểm bán (lấy file mới nhất)
    DuLieu/input_opes/     file OPES có Mã DAO (gộp TẤT CẢ file) — dùng để đối soát
    DuLieu/input_nhadong/  file "Báo cáo Nhà đồng" OPES gửi hằng tháng (lấy file mới nhất) — tỷ trọng nhà đồng

Kết quả: docs/data.enc.js — số liệu ĐÃ TỔNG HỢP và MÃ HOÁ bằng mật khẩu trong file matkhau.txt
(matkhau.txt không đẩy lên GitHub). Không xuất tên khách hàng, SĐT, biển số, tên cán bộ bán.
"""
import argparse, datetime as dt, glob, json, os, re, statistics, sys
from collections import defaultdict

try:
    import openpyxl
except ImportError:
    sys.exit("Thiếu thư viện openpyxl. Chạy: pip install openpyxl")

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import aes_ctr

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.environ.get("OCAR_DATA") or os.path.join(os.path.dirname(ROOT), "DuLieu")

CONFIG = {
    "start": "2026-09-21",
    "channel_keyword": "OPES",
    "pvi_share": 0.40,          # tỷ lệ đồng bảo hiểm của BH PVI
    "stale_days": 7,            # điểm bán quá N ngày chưa có đơn mới
}

# ---------------------------------------------------------------- tiện ích
def norm(s):
    if s is None:
        return ""
    s = str(s).replace("\n", " ")
    return re.sub(r"\s+", " ", s).strip().lower()

def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()[:10]
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None

def to_num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace(" ", "")
    try:
        return float(s)
    except ValueError:
        return 0.0

def files(sub):
    return [f for f in glob.glob(os.path.join(DATA, sub, "*.xls*"))
            if not os.path.basename(f).startswith("~$")]

def latest(sub):
    fs = files(sub)
    return max(fs, key=os.path.getmtime) if fs else None

def find_col(rows, *keys):
    keys = [norm(k) for k in keys]
    for r in rows:
        for j, v in enumerate(r):
            n = norm(v)
            if n and all(k in n for k in keys):
                return j
    return None

def file_date(path):
    name = os.path.basename(path)
    for pat, order in [(r"(?<!\d)(\d{1,2})[._\- ](\d{1,2})[._\- ](20\d{2})(?!\d)", "dmy"),
                       (r"(?<!\d)(20\d{2})[._\-](\d{1,2})[._\-](\d{1,2})(?!\d)", "ymd"),
                       (r"(?<!\d)(\d{2})(\d{2})(20\d{2})(?!\d)", "dmy")]:
        for mt in re.finditer(pat, name):
            a1, a2, a3 = (int(x) for x in mt.groups())
            try:
                return dt.date(a3, a2, a1) if order == "dmy" else dt.date(a1, a2, a3)
            except ValueError:
                pass
    return None

# ---------------------------------------------------------------- mapping
def load_mapping(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    hdr = rows[:3]
    c = dict(type=find_col(hdr, "cn/pgd"), name=find_col(hdr, "tên cn-pgd"),
             code=find_col(hdr, "branch code", "mã cn"), region=find_col(hdr, "miền"),
             prov=find_col(hdr, "tỉnh/thành phố"), vung=find_col(hdr, "vùng khcn mới"),
             vung_hhb=find_col(hdr, "vùng hhb"), pvi=find_col(hdr, "tên đơn vị pvi phụ trách"))
    if c["code"] is None or c["name"] is None:
        sys.exit("File mapping thiếu cột Branch code / Tên CN-PGD")
    g = lambda r, k: (str(r[c[k]]).strip() if c[k] is not None and c[k] < len(r) and r[c[k]] is not None else "")
    m = {}
    for r in rows[3:]:
        code = g(r, "code").upper()
        if not re.fullmatch(r"[A-Z]{2}\d{6,}", code):
            continue
        t = g(r, "type").upper(); name = g(r, "name")
        if "TTKD" in t or "OTO" in name.upper():
            kind = "HUB"
        elif "PGD" in t: kind = "PGD"
        elif "TTTC" in t: kind = "TTTC"
        elif "HHB" in t or name.upper().startswith("HH"): kind = "HHB"
        elif t == "HO": kind = "HO"
        else: kind = "CN"
        vung = g(r, "vung") or g(r, "vung_hhb")
        if kind == "HUB":
            vung = "Hub xe"
        m[code] = dict(c=code, n=name, k=kind, r=g(r, "region"), p=g(r, "prov"), vg=vung,
                       v=g(r, "pvi") or "Chưa xác định")
    return m

# ---------------------------------------------------------------- báo cáo BH PVI
def load_report(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = next((s for s in wb.worksheets if "oto" in s.title.lower() or "nvu" in s.title.lower()), wb.worksheets[0])
    rows = list(ws.iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(rows[:40]) if "số đơn bảo hiểm" in [norm(v) for v in r]
               and any("mã chi nhánh bank" in norm(v) for v in r)), None)
    if hi is None:
        sys.exit("Báo cáo BH PVI: không tìm thấy dòng tiêu đề.")
    H = [rows[hi]]
    c = dict(pol=find_col(H, "số đơn bảo hiểm"), sdbs=find_col(H, "số đơn sửa đổi bổ sung"),
             issued=find_col(H, "ngày cấp"), channel=find_col(H, "danh mục kênh bán hàng"),
             fee=find_col(H, "tổng phí bh"), bank=find_col(H, "mã chi nhánh bank"),
             partner=find_col(H, "số đơn bh từ đối tác"))
    if None in (c["pol"], c["issued"], c["fee"], c["bank"]):
        sys.exit(f"Báo cáo BH PVI thiếu cột bắt buộc: {c}")
    g = lambda r, k: (r[c[k]] if c[k] is not None and c[k] < len(r) else None)
    out = []
    for r in rows[hi + 3:]:
        pol = g(r, "pol")
        if not pol or "/" not in str(pol):
            continue
        out.append(dict(pol=str(pol).strip(), sdbs=str(g(r, "sdbs") or "").strip(),
                        issued=to_date(g(r, "issued")), channel=str(g(r, "channel") or ""),
                        fee=to_num(g(r, "fee")), bank=str(g(r, "bank") or "").strip().upper(),
                        partner=str(g(r, "partner") or "").strip().upper()))
    ob = {x["pol"]: x["bank"] for x in out if x["bank"] and not x["sdbs"]}
    op = {x["pol"]: x["partner"] for x in out if x["partner"] and not x["sdbs"]}
    for x in out:
        x["bank"] = x["bank"] or ob.get(x["pol"], "")
        x["partner"] = x["partner"] or op.get(x["pol"], "")
    return out

# ---------------------------------------------------------------- OPES (đối soát)
def load_opes_policies(path):
    """Trả về {số HĐ PL-: (ngày phát hành, mã CN, phí gồm VAT)} — chỉ dòng của BH PVI."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        hi = next((i for i, r in enumerate(rows[:20]) if any(norm(v) == "số hợp đồng bảo hiểm" for v in r)), None)
        if hi is None:
            continue
        H = [rows[hi]]
        c = dict(pol=find_col(H, "số hợp đồng bảo hiểm"), bank=find_col(H, "mã chi nhánh"),
                 issued=find_col(H, "ngày phát hành hợp đồng"), fee=find_col(H, "số phí bảo hiểm"),
                 ins=find_col(H, "nhà đồng"))
        for r in rows[hi + 1:]:
            if not r or c["pol"] >= len(r) or not r[c["pol"]]:
                continue
            if c["ins"] is not None and "PVI" not in str(r[c["ins"]] or "").upper():
                continue
            iss = r[c["issued"]] if c["issued"] is not None else None
            out[str(r[c["pol"]]).strip().upper()] = (
                iss.date() if isinstance(iss, dt.datetime) else to_date(iss),
                str(r[c["bank"]] or "").strip().upper() if c["bank"] is not None else "",
                to_num(r[c["fee"]]) if c["fee"] is not None else 0.0)
    return out

# ---------------------------------------------------------------- Nhà đồng (hằng tháng)
def load_nhadong(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        hi = next((i for i, r in enumerate(rows[:10]) if any(norm(v) == "nhà đồng" for v in r)
                   and any(norm(v) == "số hợp đồng bảo hiểm" for v in r)), None)
        if hi is None:
            continue
        H = [rows[hi]]
        c = dict(ins=find_col(H, "nhà đồng"), bank=find_col(H, "mã chi nhánh"), vung=find_col(H, "vùng"),
                 pol=find_col(H, "số hợp đồng bảo hiểm"), fee=find_col(H, "số phí bảo hiểm"),
                 fee_nv=find_col(H, "không vat"), issued=find_col(H, "ngày phát hành hợp đồng"))
        seen, out, dup = set(), [], 0
        for r in rows[hi + 1:]:
            if not r or not r[c["pol"]]:
                continue
            pol = str(r[c["pol"]]).strip().upper()
            if pol in seen:
                dup += 1
                continue
            seen.add(pol)
            iss = r[c["issued"]]
            out.append(dict(ins=str(r[c["ins"]] or "").strip().upper(), pol=pol,
                            bank=str(r[c["bank"]] or "").strip().upper(),
                            vung=str(r[c["vung"]] or "").strip() if c["vung"] is not None else "",
                            fee=to_num(r[c["fee"]]), fee_nv=to_num(r[c["fee_nv"]]) if c["fee_nv"] is not None else None,
                            issued=iss.date() if isinstance(iss, dt.datetime) else to_date(iss)))
        return out, dup
    return None, 0

# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(ROOT, "docs"))
    ap.add_argument("--plain", action="store_true", help="Ghi thêm bản KHÔNG mã hoá ra file tạm để kiểm tra (không đẩy lên)")
    a = ap.parse_args()

    rfile, mfile = latest("input"), latest("mapping")
    if not rfile: sys.exit(f"Không có báo cáo BH PVI trong {os.path.join(DATA, 'input')}")
    if not mfile: sys.exit(f"Không có file mapping trong {os.path.join(DATA, 'mapping')}")
    print(f"• Báo cáo BH PVI : {os.path.basename(rfile)}")
    print(f"• Mapping        : {os.path.basename(mfile)}")
    M = load_mapping(mfile)
    allrecs = load_report(rfile)
    start = dt.date.fromisoformat(CONFIG["start"])
    kw = CONFIG["channel_keyword"].lower()
    recs = [x for x in allrecs if x["issued"] and x["issued"] >= start and (not x["channel"] or kw in x["channel"].lower())]

    # ngày số liệu: lấy từ tên file nếu hợp lệ, không thì lấy ngày cấp mới nhất trong file
    maxd = max((x["issued"] for x in recs), default=dt.date.today())
    fd = file_date(rfile)
    rdate = fd if fd and maxd <= fd <= maxd + dt.timedelta(days=7) else maxd
    mh = re.search(r"(?<!\d)(\d{1,2})\s*[hH]\s*(\d{2})?(?![a-zA-Z])", os.path.basename(rfile))
    rtime = f"{int(mh.group(1)):02d}:{(mh.group(2) or '00')}" if mh and int(mh.group(1)) < 24 else ""

    share = CONFIG["pvi_share"]
    # ---- theo HĐ: ngày cấp đầu tiên + phí (gồm SĐBS)
    pol_first, seen = {}, set()
    daily = defaultdict(lambda: [0, 0.0])
    pts = {}
    unmapped = defaultdict(lambda: [0, 0.0])
    for x in sorted(recs, key=lambda x: x["issued"]):
        new = not x["sdbs"] and x["pol"] not in seen
        if new:
            seen.add(x["pol"])
        d = daily[x["issued"]]
        d[0] += 1 if new else 0
        d[1] += x["fee"]
        rec = M.get(x["bank"])
        if not rec:
            u = unmapped[x["bank"] or "(trống)"]; u[0] += 1 if new else 0; u[1] += x["fee"]
            continue
        z = pts.setdefault(rec["c"], dict(n=0, p=0.0, m={}, last=None, first=None, w=[0, 0.0], y=[0, 0.0]))
        z["n"] += 1 if new else 0
        if x["issued"] > rdate - dt.timedelta(days=7):          # 7 ngày gần nhất (gồm ngày số liệu)
            z["w"][0] += 1 if new else 0; z["w"][1] += x["fee"]
        if x["issued"] == rdate - dt.timedelta(days=1):         # ngày hôm trước (đủ ngày)
            z["y"][0] += 1 if new else 0; z["y"][1] += x["fee"]
        z["p"] += x["fee"]
        mk = x["issued"].strftime("%Y-%m")
        mm = z["m"].setdefault(mk, [0, 0.0]); mm[0] += 1 if new else 0; mm[1] += x["fee"]
        if new:
            iso = x["issued"].isoformat()
            z["last"] = max(z["last"] or iso, iso)
            z["first"] = min(z["first"] or iso, iso)

    days = []
    d = start
    while d <= rdate:
        n, p = daily.get(d, [0, 0.0])
        days.append([d.isoformat(), n, round(p)])
        d += dt.timedelta(days=1)

    points = []
    for code, rec in M.items():
        z = pts.get(code)
        points.append(dict(c=code, n=rec["n"], k=rec["k"], r=rec["r"], vg=rec["vg"], v=rec["v"], pr=rec["p"],
                           pn=z["n"] if z else 0, p=round(z["p"]) if z else 0,
                           m={k: [v[0], round(v[1])] for k, v in z["m"].items()} if z else {},
                           last=z["last"] if z else None, first=z["first"] if z else None,
                           w=[z["w"][0], round(z["w"][1])] if z else [0, 0], y=[z["y"][0], round(z["y"][1])] if z else [0, 0]))
    points.sort(key=lambda z: (-z["p"], z["n"]))

    total_n = len(seen)
    total_p = sum(x["fee"] for x in recs)
    data = dict(
        meta=dict(report_date=rdate.isoformat(), report_time=rtime, start=CONFIG["start"],
                  generated_at=dt.datetime.now().strftime("%Y-%m-%d %H:%M"), share=share,
                  stale_days=CONFIG["stale_days"], source=os.path.basename(rfile),
                  mapping=os.path.basename(mfile), total_n=total_n, total_p=round(total_p),
                  unmapped=[dict(c=k, n=v[0], p=round(v[1])) for k, v in unmapped.items()]),
        days=days, points=points,
    )

    # ---- đối soát OPES ↔ BH PVI
    warns = []
    ofiles = sorted(files("input_opes") + files("input_nhadong"),
                    key=lambda f: (file_date(f) or dt.date.min, os.path.getmtime(f)))
    recon = None
    if ofiles:
        O = {}
        for f in ofiles:
            O.update(load_opes_policies(f))
        omax = max((v[0] for v in O.values() if v[0]), default=None)
        pv = {}
        for x in recs:
            k = x["partner"] or x["pol"]
            e = pv.setdefault(k, dict(pol=x["pol"], d=x["issued"], b=x["bank"], fee=0.0))
            e["fee"] += x["fee"]
        name = lambda c: (M.get(c) or {}).get("n", c)
        unit = lambda c: (M.get(c) or {}).get("v", "")
        only_o = [dict(pol=k, d=v[0].isoformat() if v[0] else None, c=v[1], n=name(v[1]), v=unit(v[1]), fee=round(v[2] / 1.1))
                  for k, v in O.items() if k not in pv and v[0] and start <= v[0] < rdate]
        only_p = [dict(pol=k, d=e["d"].isoformat(), c=e["b"], n=name(e["b"]), v=unit(e["b"]), fee=round(e["fee"]))
                  for k, e in pv.items() if k not in O and omax and e["d"] <= omax]
        code_diff = [dict(pol=k, d=e["d"].isoformat(), c=e["b"], n=name(e["b"]), co=O[k][1], no=name(O[k][1]))
                     for k, e in pv.items() if k in O and O[k][1] and e["b"] and O[k][1] != e["b"]]
        fee_diff = [dict(pol=k, d=e["d"].isoformat(), c=e["b"], n=name(e["b"]), fp=round(e["fee"]), fo=round(O[k][2] / 1.1))
                    for k, e in pv.items() if k in O and O[k][2] and e["fee"]
                    and abs(O[k][2] / 1.1 - e["fee"]) > max(2000, 0.01 * e["fee"])]
        for L in (only_o, only_p, code_diff, fee_diff):
            L.sort(key=lambda z: z["d"] or "")
        recon = dict(asof=omax.isoformat() if omax else None, files=[os.path.basename(f) for f in ofiles],
                     matched=sum(1 for k in pv if k in O), pvi_total=len(pv), opes_total=len(O),
                     only_opes=only_o, only_pvi=only_p, code_diff=code_diff, fee_diff=fee_diff)
        print(f"• Đối soát OPES ({len(ofiles)} file, dữ liệu OPES đến {omax:%d/%m}): khớp {recon['matched']} HĐ | "
              f"chỉ có bên OPES {len(only_o)} | chỉ có bên BH PVI {len(only_p)} | lệch mã CN {len(code_diff)} | lệch phí {len(fee_diff)}")
    data["recon"] = recon

    # ---- tỷ trọng nhà đồng (file OPES hằng tháng)
    nfile = latest("input_nhadong")
    nd = None
    if nfile:
        rows, dup = load_nhadong(nfile)
        if rows:
            # cột "không VAT" của OPES có đang gồm VAT không? → so phần BH PVI với báo cáo BH PVI
            pvfee = {}
            for x in recs:
                k = x["partner"] or x["pol"]
                pvfee[k] = pvfee.get(k, 0.0) + x["fee"]
            ratios = [r["fee"] / pvfee[r["pol"]] for r in rows if "PVI" in r["ins"] and pvfee.get(r["pol"])]
            ratio = statistics.median(ratios) if ratios else 1.0
            div = 1.1 if abs(ratio - 1.1) < 0.03 else 1.0
            ins_name = lambda s: "BH PVI" if "PVI" in s else ("VNI (DBV)" if s in ("VNI", "DBV") else s)
            agg_m = defaultdict(lambda: [0, 0.0])
            by_v = defaultdict(lambda: defaultdict(float))
            by_u = defaultdict(lambda: defaultdict(float))
            by_c = defaultdict(lambda: defaultdict(float))
            months = set()
            for r in rows:
                if not r["issued"]:
                    continue
                ins = ins_name(r["ins"]); mk = r["issued"].strftime("%Y-%m"); months.add(mk)
                f = r["fee"] / div
                a_ = agg_m[(ins, mk)]; a_[0] += 1; a_[1] += f
                rec = M.get(r["bank"])
                by_v[(rec or {}).get("vg") or r["vung"] or "Khác"][ins] += f
                by_u[(rec or {}).get("v") or "Chưa có trong mapping"][ins] += f
                by_c[r["bank"]][ins] += f
            insurers = sorted({k[0] for k in agg_m}, key=lambda s: (s != "BH PVI", -sum(v[1] for k, v in agg_m.items() if k[0] == s)))
            pack = lambda D: [dict(k=k, t=round(sum(v.values())), s={i: round(v.get(i, 0)) for i in insurers})
                              for k, v in sorted(D.items(), key=lambda kv: -sum(kv[1].values()))]
            maxd_nd = max(r["issued"] for r in rows if r["issued"])
            nd = dict(file=os.path.basename(nfile), asof=maxd_nd.isoformat(),
                      min=min(r["issued"] for r in rows if r["issued"]).isoformat(),
                      insurers=insurers, months=sorted(months), dup=dup, vat_adjusted=div != 1.0,
                      monthly=[dict(i=k[0], m=k[1], n=v[0], p=round(v[1])) for k, v in agg_m.items()],
                      by_vung=pack(by_v), by_unit=pack(by_u),
                      by_point=[dict(c=k, n=(M.get(k) or {}).get("n", k), v=(M.get(k) or {}).get("v", ""),
                                     t=round(sum(v.values())), pvi=round(v.get("BH PVI", 0)))
                                for k, v in by_c.items()])
            print(f"• Nhà đồng: {os.path.basename(nfile)} ({len(rows)} HĐ, loại {dup} dòng trùng, dữ liệu đến {maxd_nd:%d/%m})"
                  + (" — cột phí của OPES đang GỒM VAT, đã quy về chưa VAT (÷1,1)" if div != 1.0 else ""))
    data["nhadong"] = nd

    # ---- kiểm tra
    if unmapped:
        warns.append(f"Mã chi nhánh chưa có trong mapping: " + ", ".join(f"{k} ({v[0]} HĐ)" for k, v in unmapped.items()))
    fut = [x for x in recs if x["issued"] > rdate]
    if fut:
        warns.append(f"{len(fut)} dòng có ngày cấp sau ngày báo cáo {rdate:%d/%m}")

    # ---- mã hoá & ghi
    pw_file = os.path.join(ROOT, "matkhau.txt")
    pw = open(pw_file, encoding="utf-8").read().strip() if os.path.exists(pw_file) else ""
    if not pw:
        sys.exit("Chưa có mật khẩu: tạo file matkhau.txt trong thư mục Monitor, ghi mật khẩu vào dòng đầu tiên rồi chạy lại.")
    raw = json.dumps(data, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    enc = aes_ctr.encrypt(raw, pw)
    enc["hint"] = f"{rdate:%d/%m/%Y}"
    os.makedirs(a.out, exist_ok=True)
    with open(os.path.join(a.out, "data.enc.js"), "w", encoding="utf-8") as f:
        f.write("window.OCAR_ENC=" + json.dumps(enc) + ";\n")
    idx = os.path.join(a.out, "index.html")
    if os.path.exists(idx):
        html = open(idx, encoding="utf-8").read()
        html2 = re.sub(r'src="data\.enc\.js(?:\?v=[^"]*)?"', f'src="data.enc.js?v={dt.datetime.now():%Y%m%d%H%M%S}"', html)
        if html2 != html:
            open(idx, "w", encoding="utf-8", newline="").write(html2)
    if a.plain:
        with open(os.path.join(ROOT, "data_kiemtra_KHONG_DAY_LEN.json"), "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=1)

    print(f"• Ngày số liệu: {rdate:%d/%m/%Y}{(' ' + rtime) if rtime else ''}")
    print(f"• Luỹ kế từ {start:%d/%m}: {total_n} HĐBH | {total_p:,.0f}đ (100%, chưa VAT) | phần BH PVI {share:.0%}: {total_p*share:,.0f}đ")
    act = sum(1 for z in points if z["pn"])
    print(f"• Điểm bán có đơn: {act}/{len(points)}")
    for w in warns:
        print(f"  ⚠ {w}")
    print("✔ Đã ghi docs/data.enc.js (đã mã hoá)")

if __name__ == "__main__":
    main()
